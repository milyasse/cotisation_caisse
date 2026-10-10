"""
Génère les canevas Excel d'import de l'application Caisse commune :
  - caisse-donnees-reelles.xlsx : pré-rempli avec l'historique du groupe (transcrit des captures du 10/10/2026)
  - caisse-modele-vierge.xlsx   : même structure, une ligne d'exemple par feuille
Usage : python3 scripts/canevas.py <dossier de sortie>
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
MEMBRES = ["امحمد", "بوعزة", "بنعاشير", "موسى"]
CATEGORIES = ["Paiement du personnel", "Charges communes", "Achat de matériel", "Entretien et réparation", "Autres dépenses"]
MODES = ["Espèces", "Virement", "Autre"]
MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]

F = "Arial"
HEAD = PatternFill("solid", fgColor="17633D")
INPUT = PatternFill("solid", fgColor="FFF8D6")
THIN = Side(style="thin", color="D3DBD5")

# ---------------------------------------------------------------- Données (feuille « المساهمات »)
# (année, mois, date de réception, [امحمد, بوعزة, بنعاشير, موسى], dépenses du mois)
ROWS = [
    (2025, 3, "01/03/2025", [450, 400, 400, 400], 1600),
    (2025, 4, "16/03/2025", [400, 400, 400, 400], 1850),
    (2025, 5, "26/04/2025", [450, 400, 400, 400], 1450),
    (2025, 6, "31/05/2025", [450, 400, 400, 400], 1600),
    (2025, 7, "30/06/2025", [400, 400, 400, 400], 1750),  # « 31/06/2025 » dans le fichier d'origine
    (2025, 8, "20/07/2025", [400, 400, 400, 400], 1400),
    (2025, 9, "30/08/2025", [400, 400, 400, 400], 1750),
    (2025, 10, "27/09/2025", [400, 400, 400, 400], 1400),
    (2025, 11, "01/11/2025", [400, 400, 400, 400], 1400),
    (2025, 12, "28/11/2025", [400, 400, 400, 400], 1750),
    (2026, 1, "01/01/2026", [400, 400, 400, 400], 1750),
    (2026, 2, "01/02/2026", [400, 400, 400, 400], 1400),
    (2026, 3, "01/03/2026", [400, 400, 400, 400], 1400),
    (2026, 4, "01/04/2026", [400, 400, 400, 400], 1400),
    (2026, 5, "02/05/2026", [400, 400, 400, 400], 1750),
    (2026, 6, "01/06/2026", [400, 400, 400, 400], 1400),
    (2026, 7, "01/07/2026", [400, 400, 400, 400], 1400),
    (2026, 8, "01/08/2026", [400, 400, 400, 400], 1750),
    (2026, 9, "01/09/2026", [600, 600, None, 400], 1400),
    (2026, 10, "01/10/2026", [600, None, None, 400], 350),
]
LAST_DAY = {1: 31, 2: 28, 3: 31, 4: 30, 5: 31, 6: 30, 7: 31, 8: 31, 9: 30, 10: 31, 11: 30, 12: 31}
TODAY = "10/10/2026"


def style_header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def finish(ws, ncols, nrows, money_col, extra_rows=200):
    """Mise en forme des lignes à remplir (fond jaune clair = cellules à saisir)."""
    for r in range(2, nrows + extra_rows + 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=F)
            cell.fill = INPUT
            cell.border = Border(bottom=THIN)
            cell.number_format = "#,##0.00" if c == money_col else "@"
            if c != money_col:
                cell.value = None if cell.value is None else str(cell.value)


def versements_sheet(wb, rows):
    ws = wb.active
    ws.title = "Versements"
    style_header(ws, ["Date", "Membre", "Mois concerné", "Montant (MAD)", "Commentaire"], [14, 18, 16, 16, 40])
    for date, membre, mois, montant, com in rows:
        ws.append([date, membre, mois, montant, com])
    finish(ws, 5, len(rows), money_col=4)
    dv = DataValidation(type="list", formula1='"' + ",".join(MEMBRES) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("B2:B500")
    ws["A1"].comment = Comment("Date réelle où l'argent a été reçu, au format JJ/MM/AAAA.", "Caisse commune")
    ws["C1"].comment = Comment("Mois que paie ce versement, au format MM/AAAA (ex. 04/2025). Peut être différent du mois de la date.", "Caisse commune")
    ws["B1"].comment = Comment("Nom exactement comme dans l'application. Un nom inconnu crée un nouveau membre.", "Caisse commune")
    return ws


def depenses_sheet(wb, rows):
    ws = wb.create_sheet("Dépenses")
    style_header(ws, ["Date", "Catégorie", "Description", "Bénéficiaire", "Montant (MAD)", "Mode", "Référence", "Commentaire"], [14, 22, 36, 20, 16, 12, 14, 30])
    for r in rows:
        ws.append(list(r))
    finish(ws, 8, len(rows), money_col=5)
    dv1 = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True, showErrorMessage=False)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(MODES) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("B2:B500")
    dv2.add("F2:F500")
    ws["A1"].comment = Comment("Date réelle de la dépense, au format JJ/MM/AAAA.", "Caisse commune")
    return ws


def mode_emploi(wb, prefilled):
    ws = wb.create_sheet("Mode d'emploi")
    ws.column_dimensions["A"].width = 100
    lines = [
        ("Caisse commune — canevas d'import", True),
        ("", False),
        ("Cellules à remplir : les cellules sur fond jaune clair des feuilles « Versements » et « Dépenses ». Ne changez pas la ligne 1 (titres).", False),
        ("", False),
        ("Feuille « Versements » — une ligne par versement reçu :", True),
        ("  Date : date réelle de réception, JJ/MM/AAAA (ex. 01/03/2025)", False),
        ("  Membre : nom du membre (liste déroulante). Un nouveau nom crée un nouveau membre.", False),
        ("  Mois concerné : mois payé, MM/AAAA (ex. 04/2025). Un versement de mars pour avril : Date 16/03/2025, Mois 04/2025.", False),
        ("  Montant (MAD) : nombre, ex. 400 ou 450,50. Plusieurs versements pour un même mois = plusieurs lignes.", False),
        ("", False),
        ("Feuille « Dépenses » — une ligne par dépense (ou un total par mois) :", True),
        ("  Date (JJ/MM/AAAA), Catégorie (liste), Description, Bénéficiaire, Montant (MAD), Mode (liste), Référence et Commentaire facultatifs.", False),
        ("", False),
        ("Importer dans l'application (iPhone) :", True),
        ("  1. Enregistrez ce fichier dans « Fichiers » (depuis Google Sheets : ⋮ › Partager et exporter › Enregistrer sous › Excel).", False),
        ("  2. Application › Paramètres › Importer l'ancien fichier Excel › choisissez ce fichier.", False),
        ("  3. Feuille « Versements » › Contrôler les données › Importer. (La 1re fois, laissez cochée « Remplacer les données de démonstration ».)", False),
        ("  4. Recommencez avec la feuille « Dépenses » › Contrôler les données › Importer.", False),
        ("  Rien n'est jamais importé deux fois : un réimport du même fichier n'ajoute que les nouvelles lignes.", False),
    ]
    if prefilled:
        lines += [
            ("", False),
            ("Contenu pré-rempli (à vérifier avant import) :", True),
            ("  Versements de mars 2025 à octobre 2026, recopiés de la feuille « المساهمات » (77 lignes).", False),
            ("  Dépenses : le détail de la feuille « المصاريف » — 84 salaires hebdomadaires du gardien et 6 recharges d'électricité (90 lignes).", False),
            ("  Correction : semaine du 08/09 au 14/09 payée « 06/08/2025 » dans le fichier, remplacé par 06/09/2025 (à vérifier).", False),
            ("  Correction : « 31/06/2025 » (date inexistante) remplacé par 30/06/2025.", False),
            ("  Résultat attendu après les deux imports : solde en caisse 1 600,00 MAD au 10/10/2026, comme dans le fichier d'origine.", False),
        ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=F, bold=bold, size=13 if i == 1 else 11, color="17633D" if bold else "15211A")
        c.alignment = Alignment(wrap_text=True, vertical="top")


def depenses_detaillees():
    """Lit l'extrait réel de la feuille « المصاريف » (tests/fixtures/extrait-depenses-reel.tsv)."""
    import csv, os
    path = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures", "extrait-depenses-reel.tsv")
    rows = list(csv.reader(open(path, encoding="utf-8"), delimiter="\t"))
    deps, section = [], ""
    for r in rows[3:]:
        r = r + [""] * 30
        week, rub, date, payer = r[0].strip(), r[1].strip(), r[2].strip(), r[3].strip()
        if week == "الأسبوع":  # en-tête répété
            continue
        amounts = [x.strip() for x in r[5:27] if x.strip() not in ("", "-")]
        if not amounts:
            if not date and rub and week in ("العدد", ""):
                section = rub
            continue
        if not date or "مجموع" in week + rub:
            continue
        amount = float(amounts[0].replace(" ", "").replace(",", "."))
        label = rub or section
        cat = "Paiement du personnel" if "أجرة" in label or "الحارس" in label else "Charges communes"
        benef = "Gardien" if cat == "Paiement du personnel" else ""
        comment = f"Payé par {payer}" if payer else None
        if date == "06/08/2025" and "08/09" in rub:
            date = "06/09/2025"
            comment = (comment or "") + " · date corrigée (fichier : 06/08/2025)"
        deps.append((date, cat, label, benef, amount, "Espèces", f"Semaine {week}" if week.isdigit() and cat == "Paiement du personnel" else None, comment))
    return deps


def build(prefilled: bool, path: str):
    wb = Workbook()
    if prefilled:
        pays, deps = [], []
        for year, month, date, amounts, dep in ROWS:
            for name, a in zip(MEMBRES, amounts):
                if a:
                    pays.append((date, name, f"{month:02d}/{year}", a, None))
        deps = depenses_detaillees()
    else:
        pays = [("01/10/2026", MEMBRES[0], "10/2026", 400, "Exemple — à remplacer")]
        deps = [("05/10/2026", "Paiement du personnel", "Salaire hebdomadaire du gardien", "Gardien", 350, "Espèces", None, "Exemple — à remplacer")]
    versements_sheet(wb, pays)
    depenses_sheet(wb, deps)
    mode_emploi(wb, prefilled)
    wb.save(path)
    print(path, len(pays), "versements,", len(deps), "dépenses")


build(True, f"{OUT}/caisse-donnees-reelles.xlsx")
build(False, f"{OUT}/caisse-modele-vierge.xlsx")
